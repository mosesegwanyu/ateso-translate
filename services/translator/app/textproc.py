"""Document → paragraphs → sentences, and language detection. No ML; fully unit-tested."""
import json,re,io,os
ABBR={'mr','mrs','ms','dr','prof','st','vs','etc','e.g','i.e','no','fig','rev','hon','eng','sr','jr','mt'}
_WL=json.load(open(os.path.join(os.path.dirname(__file__),'wordlists.json'),encoding='utf8'))
_SETS={k:set(v) for k,v in _WL.items()}
tok=lambda s:[w.strip("'") for w in re.findall(r"[\w']+",s.lower().replace('’',"'")) if w.strip("'")]

def detect(text):
    """Return (lang, confidence 0-1). 'unknown' if too little evidence."""
    w=tok(text); e=sum(x in _SETS['eng'] for x in w); t=sum(x in _SETS['teo'] for x in w)
    if e+t<2: return 'unknown',0.0
    return ('eng' if e>=t else 'teo'), round(max(e,t)/(e+t),2)

def split_sentences(par,max_words=60):
    par=re.sub(r'\s+',' ',par).strip()
    if not par: return []
    out,start=[],0
    for m in re.finditer(r'[.!?…]+["”’)\]]*\s+(?=["“‘(\[]?[A-Z0-9])',par):
        head=par[start:m.end()].strip(); last=re.findall(r"[\w.]+",head)[-1].lower().rstrip('.') if re.findall(r"[\w.]+",head) else ''
        if last in ABBR or re.fullmatch(r'[a-z]',last): continue   # Dr. Smith, initials
        out.append(head); start=m.end()
    if par[start:].strip(): out.append(par[start:].strip())
    final=[]
    for s in out:   # break over-long sentences at ; or , so the model isn't fed essays
        if len(s.split())<=max_words: final.append(s); continue
        buf=''
        for piece in re.split(r'(?<=[;:,])\s+',s):
            if buf and len((buf+' '+piece).split())>max_words: final.append(buf); buf=piece
            else: buf=(buf+' '+piece).strip()
        if buf: final.append(buf)
    return final

def paragraphs_from_text(text):
    blocks=re.split(r'\n\s*\n',text.replace('\r\n','\n'))
    return [re.sub(r'\s*\n\s*',' ',b).strip() for b in blocks if b.strip()]

def extract(filename,data):
    n=filename.lower()
    if n.endswith('.txt') or n.endswith('.md'): return paragraphs_from_text(data.decode('utf-8','replace'))
    if n.endswith('.docx'):
        import docx; d=docx.Document(io.BytesIO(data)); ps=[p.text.strip() for p in d.paragraphs if p.text.strip()]
        for t in d.tables:
            for r in t.rows:
                for c in r.cells:
                    if c.text.strip(): ps.append(c.text.strip())
        return ps
    if n.endswith('.pdf'):
        import pypdf; r=pypdf.PdfReader(io.BytesIO(data))
        if len(r.pages)>60: raise ValueError('PDF has more than 60 pages')
        text='\n\n'.join((p.extract_text() or '') for p in r.pages)
        if len(text.strip())<20: raise ValueError('No text found. This PDF may be a scan; scanned PDFs need OCR, which is not supported yet.')
        text=re.sub(r'(?<![.!?:\n])\n(?!\n)',' ',text)   # undo hard line wraps
        return paragraphs_from_text(text)
    raise ValueError('Unsupported file type. Use .txt, .docx or .pdf')
