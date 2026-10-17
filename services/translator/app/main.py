import os,io,threading
from fastapi import FastAPI,UploadFile,File,HTTPException,Header,Response
from fastapi.middleware.cors import CORSMiddleware
from pydantic import BaseModel
from . import textproc
if os.getenv('ENGINE','sunbird')=='local': from . import engine_nllb as engine
else: from . import engine_sunbird as engine
app=FastAPI(title='Ateso Translate service'); TOKEN=os.getenv('API_TOKEN','')
app.add_middleware(CORSMiddleware,allow_origins=[o for o in os.getenv('ALLOWED_ORIGINS','*').split(',')],allow_methods=['*'],allow_headers=['*'])
LANGS={'eng','teo'}
def guard(auth):
    if TOKEN and auth!=f'Bearer {TOKEN}': raise HTTPException(401,'Unauthorized')
@app.on_event('startup')
def boot(): threading.Thread(target=engine.load,daemon=True).start()   # server answers /health while the model loads
@app.get('/health')
def health(): return {'ready':engine.ready(),'model':engine.MODEL_ID}
class Det(BaseModel): text:str
@app.post('/detect')
def detect(b:Det,authorization:str=Header('')):
    guard(authorization); l,c=textproc.detect(b.text[:5000]); return {'lang':l,'confidence':c}
class Tr(BaseModel): sentences:list[str]; src:str; tgt:str
@app.post('/translate')
def translate(b:Tr,authorization:str=Header('')):
    guard(authorization)
    if b.src not in LANGS or b.tgt not in LANGS or b.src==b.tgt: raise HTTPException(400,'src/tgt must be different, from eng|teo')
    if len(b.sentences)>16 or any(len(s)>1500 for s in b.sentences): raise HTTPException(400,'Too many or too long sentences')
    if not engine.ready(): raise HTTPException(503,'Translation engine is not configured or still loading')
    try: r=engine.translate([s.strip() for s in b.sentences],b.src,b.tgt)
    except RuntimeError as e: raise HTTPException(429 if type(e).__name__=='RateLimited' else 502,str(e))
    except Exception: raise HTTPException(502,'Translation engine error')
    return {'model':engine.MODEL_ID,'results':[{'text':t,'confidence':c} for t,c in r]}
@app.post('/extract')
async def extract(file:UploadFile=File(...),authorization:str=Header('')):
    guard(authorization); data=await file.read()
    if len(data)>8_000_000: raise HTTPException(413,'File larger than 8 MB')
    try: ps=textproc.extract(file.filename or '',data)
    except ValueError as e: raise HTTPException(400,str(e))
    except Exception: raise HTTPException(400,'Could not read this file')
    out=[{'text':p,'sentences':textproc.split_sentences(p)} for p in ps]
    words=sum(len(p.split()) for p in ps); l,c=textproc.detect(' '.join(ps)[:20000])
    return {'paragraphs':out,'words':words,'sentences':sum(len(p['sentences']) for p in out),'lang':l,'confidence':c}
class Seg(BaseModel): src:str=''; tgt:str=''
class Exp(BaseModel): title:str='Translation'; format:str='docx'; paragraphs:list[Seg]; bilingual:bool=False
@app.post('/export')
def export(b:Exp,authorization:str=Header('')):
    guard(authorization)
    if b.format=='txt':
        body='\n\n'.join((f'{p.src}\n{p.tgt}' if b.bilingual else p.tgt) for p in b.paragraphs)
        return Response(body.encode('utf-8'),media_type='text/plain; charset=utf-8',headers={'Content-Disposition':'attachment; filename="translation.txt"'})
    import docx; d=docx.Document(); d.add_heading(b.title[:120],1)
    for p in b.paragraphs:
        if b.bilingual: d.add_paragraph(p.src).italic=True
        d.add_paragraph(p.tgt)
    d.add_paragraph('Machine translation draft. Have a fluent speaker review it before relying on it.').italic=True
    bio=io.BytesIO(); d.save(bio)
    return Response(bio.getvalue(),media_type='application/vnd.openxmlformats-officedocument.wordprocessingml.document',headers={'Content-Disposition':'attachment; filename="translation.docx"'})
