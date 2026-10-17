import sys,os,io; sys.path.insert(0,os.path.join(os.path.dirname(__file__),'..'))
from app import textproc as t
def test_split_basic(): assert t.split_sentences('I went home. She stayed! Why? Okay.')==['I went home.','She stayed!','Why?','Okay.']
def test_abbrev(): assert t.split_sentences('Dr. Okello met Mr. Ojok. They talked.')==['Dr. Okello met Mr. Ojok.','They talked.']
def test_long_split():
    s=', '.join(['word '*10]*10).strip(); parts=t.split_sentences(s,max_words=30); assert len(parts)>1 and all(len(p.split())<=35 for p in parts)
def test_detect(): 
    assert t.detect('The children are going to school and the people are happy')[0]=='eng'
    assert t.detect('Akoto eong alosit oibuga owiken keda itunga')[0]=='teo'
    assert t.detect('hmm')[0]=='unknown'
def test_txt_extract(): assert t.extract('a.txt',b'One\ntwo\n\nThree')==['One two','Three']
def test_docx_roundtrip():
    import docx; d=docx.Document(); d.add_paragraph('Hello world.'); d.add_paragraph('Second one.'); b=io.BytesIO(); d.save(b)
    assert t.extract('x.docx',b.getvalue())==['Hello world.','Second one.']
def test_bad_type():
    try: t.extract('x.exe',b''); assert False
    except ValueError: pass
